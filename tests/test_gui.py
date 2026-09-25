"""Drives the real Tk widgets; skipped where Tk or a display is unavailable."""

import time

import openpyxl
import pytest

tk = pytest.importorskip("tkinter")


@pytest.fixture
def app(monkeypatch):
    from tkinter import messagebox

    monkeypatch.setattr(messagebox, "askyesno", lambda *a, **k: True)
    try:
        from excel_mass_replacer.gui import App

        instance = App()
    except tk.TclError as exc:
        pytest.skip(f"no display available: {exc}")
    yield instance
    instance.destroy()


def _pump(app, seconds=10.0):
    deadline = time.time() + seconds
    while time.time() < deadline:
        app.update()
        if not app._busy:
            return
        time.sleep(0.02)
    raise AssertionError("the run never finished")


@pytest.fixture
def one_file(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Acme Holdings"
    wb.save(tmp_path / "demo.xlsx")
    return tmp_path


def test_preview_does_not_write(app, one_file):
    app.folder.set(str(one_file))
    app.find_text.set("Acme")
    app.replace_text.set("Globex")
    app.preview_button.invoke()
    _pump(app)
    assert openpyxl.load_workbook(one_file / "demo.xlsx").active["A1"].value == (
        "Acme Holdings"
    )
    assert "would change" in app.log.get("1.0", "end")


def test_apply_writes_and_backs_up_from_worker_thread(app, one_file):
    """Regression: settings must be read on the main thread, not in the worker."""
    app.folder.set(str(one_file))
    app.find_text.set("Acme")
    app.replace_text.set("Globex")
    app.apply_button.invoke()
    _pump(app)
    log = app.log.get("1.0", "end")
    assert "FAILED" not in log and "ERROR" not in log
    assert openpyxl.load_workbook(one_file / "demo.xlsx").active["A1"].value == (
        "Globex Holdings"
    )
    assert (one_file / "demo.xlsx.bak").exists()


def test_buttons_are_re_enabled_after_a_run(app, one_file):
    app.folder.set(str(one_file))
    app.find_text.set("Acme")
    app.replace_text.set("Globex")
    app.preview_button.invoke()
    _pump(app)
    assert "disabled" not in app.preview_button.state()
    assert "disabled" not in app.apply_button.state()

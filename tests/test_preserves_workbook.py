"""Guards for everything a replacement must leave alone.

These exist because the tool used to rebuild each workbook through openpyxl.
The text came out right, and on the way out the file quietly lost the results
Excel had cached for every formula, its images, its printer settings and its
form controls. A program reading such a file afterwards sees blank cells, so
the damage stayed invisible until something downstream could not read the data.
"""

import shutil
import zipfile

import openpyxl
import pytest
from openpyxl.worksheet.datavalidation import DataValidation

from excel_mass_replacer.core import (
    BACKUP_DIR_NAME,
    Rule,
    discover_files,
    process_file,
    run,
)

INLINE_LIST = '"PT Lama,PT Tengah,NON PT"'
RANGE_LIST = "=Daftar!$A$1:$A$2"
EXTRA_PARTS = {
    "xl/media/image1.png": b"\x89PNG\r\n\x1a\n-pretend-this-is-a-logo",
    "xl/printerSettings/printerSettings1.bin": b"\x00\x01printer settings\x02",
    "xl/ctrlProps/ctrlProp1.xml": b"<formControlPr xmlns='x' objectType='CheckBox'/>",
}


def _add_cached_value_and_extras(path):
    """Make the file look like one Excel itself saved.

    openpyxl never writes a cached formula result, so a workbook it created has
    nothing to lose. Excel always writes one, and that is the value every other
    program reads. The extra parts stand in for the images, printer settings
    and form controls a real workbook carries.
    """
    with zipfile.ZipFile(path) as archive:
        items = [(info, archive.read(info.filename)) for info in archive.infolist()]

    with zipfile.ZipFile(path, "w") as archive:
        for info, data in items:
            if info.filename == "xl/worksheets/sheet1.xml":
                text = data.decode("utf-8")
                assert "<f>A2+A3</f>" in text, "fixture no longer matches openpyxl"
                text = text.replace("<f>A2+A3</f>", "<f>A2+A3</f><v>111</v>")
                data = text.encode("utf-8")
            archive.writestr(info, data)
        for name, blob in EXTRA_PARTS.items():
            archive.writestr(name, blob)


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "po.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PO"
    ws["A1"] = "PT Lama"
    ws["A2"] = 100
    ws["A3"] = 11
    ws["A4"] = "=A2+A3"
    ws["B1"] = "sent to PT Lama today"

    typed = DataValidation(type="list", formula1=INLINE_LIST, allowBlank=True)
    ws.add_data_validation(typed)
    typed.add(ws["C1"])
    ws["C1"] = "PT Lama"

    listed = DataValidation(type="list", formula1=RANGE_LIST, allowBlank=True)
    ws.add_data_validation(listed)
    listed.add(ws["D1"])
    ws["D1"] = "PT Lama"

    other = wb.create_sheet("Daftar")
    other["A1"] = "PT Lama"
    other["A2"] = "PT Tengah"

    wb.save(path)
    _add_cached_value_and_extras(path)
    return path


def _parts(path):
    with zipfile.ZipFile(path) as archive:
        return {info.filename: archive.read(info.filename) for info in archive.infolist()}


def _validations(path, sheet="PO"):
    ws = openpyxl.load_workbook(path)[sheet]
    return [str(dv.formula1) for dv in ws.data_validations.dataValidation]


def test_cached_formula_result_survives(workbook):
    """The whole point: a program reading the file still sees the computed value."""
    before = openpyxl.load_workbook(workbook, data_only=True)["PO"]["A4"].value
    assert before == 111

    process_file(workbook, [Rule("PT Lama", "PT Baru")], apply=True, backup=False)

    after = openpyxl.load_workbook(workbook, data_only=True)["PO"]["A4"].value
    assert after == 111, "the cached result was dropped, so readers see a blank cell"
    assert openpyxl.load_workbook(workbook)["PO"]["A4"].value == "=A2+A3"


def test_images_and_settings_survive(workbook):
    before = _parts(workbook)
    process_file(workbook, [Rule("PT Lama", "PT Baru")], apply=True, backup=False)
    after = _parts(workbook)

    for name, blob in EXTRA_PARTS.items():
        assert name in after, f"{name} was dropped"
        assert after[name] == blob, f"{name} was altered"
    assert set(before) == set(after), "the package lost or gained parts"


def test_untouched_parts_are_byte_identical(workbook):
    """Only the parts that actually hold the replaced text may be rewritten."""
    before = _parts(workbook)
    process_file(workbook, [Rule("PT Lama", "PT Baru")], apply=True, backup=False)
    after = _parts(workbook)

    changed = {name for name in before if before[name] != after[name]}
    carries_text = {"xl/sharedStrings.xml"} | {
        name for name in before if name.startswith("xl/worksheets/")
    }
    assert changed <= carries_text, (
        f"parts with no text to replace were rewritten anyway: "
        f"{sorted(changed - carries_text)}"
    )
    assert changed, "the fixture should have had something to replace"


def test_typed_in_dropdown_choices_are_replaced(workbook):
    """Otherwise a cell holds a value its own dropdown no longer offers."""
    assert INLINE_LIST in _validations(workbook)

    result = process_file(
        workbook, [Rule("PT Lama", "PT Baru")], apply=True, backup=False
    )

    lists = _validations(workbook)
    assert '"PT Baru,PT Tengah,NON PT"' in lists
    assert not any("PT Lama" in item for item in lists)
    assert result.dropdown_lists == 1, "a changed validation rule must be reported"


def test_dropdown_reading_from_another_sheet_follows_its_cells(workbook):
    """A list pointing at cells needs no special handling: the cells carry the text."""
    process_file(workbook, [Rule("PT Lama", "PT Baru")], apply=True, backup=False)

    assert RANGE_LIST in _validations(workbook), "the reference must stay a reference"
    assert openpyxl.load_workbook(workbook)["Daftar"]["A1"].value == "PT Baru"


def test_dropdown_count_is_zero_when_only_cells_change(workbook):
    result = process_file(
        workbook, [Rule("sent to", "delivered to")], apply=True, backup=False
    )
    assert result.replacements == 1
    assert result.dropdown_lists == 0


def test_backup_is_a_real_workbook_in_one_folder(workbook):
    result = process_file(workbook, [Rule("PT Lama", "PT Baru")], apply=True)

    assert result.backup.exists()
    assert result.backup.suffix == ".xlsx", "the copy must stay a workbook"
    assert result.backup.name == workbook.name
    assert BACKUP_DIR_NAME in result.backup.parts
    assert not list(workbook.parent.glob("*.bak"))

    kept = openpyxl.load_workbook(result.backup, data_only=True)["PO"]
    assert kept["A1"].value == "PT Lama", "the copy must predate the replacement"
    assert kept["A4"].value == 111


def test_one_run_puts_every_backup_in_the_same_folder(tmp_path, workbook):
    nested = tmp_path / "sub"
    nested.mkdir()
    shutil.copy2(workbook, nested / "po2.xlsx")

    files = discover_files(tmp_path)
    summary = run(files, [Rule("PT Lama", "PT Baru")], apply=True)

    assert summary.backup_dir is not None
    copies = sorted(p.name for p in summary.backup_dir.rglob("*.xlsx"))
    assert copies == ["po.xlsx", "po2.xlsx"]
    # the sub-folder is mirrored, so two files of the same name cannot collide
    assert (summary.backup_dir / "sub" / "po2.xlsx").exists()


def test_a_later_run_ignores_the_backup_folder(tmp_path, workbook):
    run(discover_files(tmp_path), [Rule("PT Lama", "PT Baru")], apply=True)

    second = discover_files(tmp_path)
    assert [p.name for p in second] == ["po.xlsx"], (
        "backups were picked up as input, so they would be rewritten too"
    )


def test_running_twice_leaves_the_cached_result_alone(workbook):
    rules = [Rule("PT Lama", "PT Baru")]
    process_file(workbook, rules, apply=True, backup=False)
    second = process_file(workbook, rules, apply=True, backup=False)

    assert second.replacements == 0
    assert openpyxl.load_workbook(workbook, data_only=True)["PO"]["A4"].value == 111

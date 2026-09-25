import pytest


@pytest.fixture
def sample_dir(tmp_path):
    import openpyxl
    import xlwt

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data Acme"
    ws["A1"] = "Acme Holdings"
    ws["A2"] = "invoice for Acme"
    ws["A3"] = 12345
    ws["A4"] = None
    ws["A5"] = "=CONCATENATE(\"Acme\",A1)"
    ws2 = wb.create_sheet("Sheet2")
    ws2["B2"] = "acme in lowercase"
    wb.save(tmp_path / "book1.xlsx")

    wb2 = openpyxl.Workbook()
    wb2.active["A1"] = "nothing to see here"
    wb2.save(tmp_path / "book2.xlsx")

    nested = tmp_path / "sub"
    nested.mkdir()
    wb3 = openpyxl.Workbook()
    wb3.active["A1"] = "Acme nested"
    wb3.save(nested / "book3.xlsx")

    old = xlwt.Workbook()
    sheet = old.add_sheet("Legacy")
    sheet.write(0, 0, "Acme legacy")
    sheet.write(1, 0, 42)
    old.save(str(tmp_path / "legacy.xls"))

    (tmp_path / "~$book1.xlsx").write_bytes(b"lock file")
    (tmp_path / "notes.txt").write_text("Acme")
    return tmp_path

import openpyxl
import pytest
import xlrd

from excel_mass_replacer.core import Rule, discover_files, process_file, run


def test_discover_skips_lock_and_non_excel(sample_dir):
    names = sorted(p.name for p in discover_files(sample_dir))
    assert names == ["book1.xlsx", "book2.xlsx", "book3.xlsx", "legacy.xls"]


def test_discover_non_recursive(sample_dir):
    names = sorted(p.name for p in discover_files(sample_dir, recursive=False))
    assert names == ["book1.xlsx", "book2.xlsx", "legacy.xls"]


def test_dry_run_changes_nothing(sample_dir):
    before = (sample_dir / "book1.xlsx").read_bytes()
    result = process_file(sample_dir / "book1.xlsx", [Rule("PT Lama", "PT Baru")])
    assert result.replacements == 2
    assert result.written is False
    assert (sample_dir / "book1.xlsx").read_bytes() == before


def test_apply_writes_and_backs_up(sample_dir):
    path = sample_dir / "book1.xlsx"
    result = process_file(path, [Rule("PT Lama", "PT Baru")], apply=True)
    assert result.written and result.backup.exists()
    ws = openpyxl.load_workbook(path).active
    assert ws["A1"].value == "PT Baru Jaya"
    assert ws["A2"].value == "invoice for PT Baru"
    assert ws["A3"].value == 12345
    assert ws["A5"].value == '=CONCATENATE("PT Lama",A1)'


def test_ignore_case(sample_dir):
    result = process_file(
        sample_dir / "book1.xlsx", [Rule("pt lama", "X", match_case=False)]
    )
    assert result.replacements == 3


def test_whole_cell(sample_dir):
    result = process_file(
        sample_dir / "book1.xlsx", [Rule("PT Lama Jaya", "X", whole_cell=True)]
    )
    assert result.replacements == 1


def test_include_formulas(sample_dir):
    path = sample_dir / "book1.xlsx"
    process_file(path, [Rule("PT Lama", "PT Baru")], apply=True, include_formulas=True)
    assert openpyxl.load_workbook(path).active["A5"].value == (
        '=CONCATENATE("PT Baru",A1)'
    )


def test_sheet_names(sample_dir):
    path = sample_dir / "book1.xlsx"
    process_file(path, [Rule("Lama", "Baru")], apply=True, include_sheet_names=True)
    assert "Data PT Baru" in openpyxl.load_workbook(path).sheetnames


def test_legacy_xls(sample_dir):
    path = sample_dir / "legacy.xls"
    result = process_file(path, [Rule("PT Lama", "PT Baru")], apply=True)
    assert result.written
    sheet = xlrd.open_workbook(path).sheet_by_index(0)
    assert sheet.cell_value(0, 0) == "PT Baru legacy"
    assert sheet.cell_value(1, 0) == 42


def test_replacement_is_literal_not_regex(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active["A1"] = "a.b"
    wb.save(tmp_path / "x.xlsx")
    process_file(tmp_path / "x.xlsx", [Rule("a.b", r"c\1d$&")], apply=True)
    assert openpyxl.load_workbook(tmp_path / "x.xlsx").active["A1"].value == r"c\1d$&"


def test_literal_find_does_not_match_regex(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active["A1"] = "axb"
    wb.save(tmp_path / "x.xlsx")
    assert process_file(tmp_path / "x.xlsx", [Rule("a.b", "Z")]).replacements == 0


def test_regex_mode(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active["A1"] = "INV-2024-001"
    wb.save(tmp_path / "x.xlsx")
    process_file(
        tmp_path / "x.xlsx", [Rule(r"INV-(\d{4})", r"BILL-\1", regex=True)], apply=True
    )
    assert openpyxl.load_workbook(tmp_path / "x.xlsx").active["A1"].value == (
        "BILL-2024-001"
    )


def test_corrupt_file_does_not_stop_the_batch(sample_dir):
    (sample_dir / "broken.xlsx").write_bytes(b"this is not a zip archive")
    summary = run(discover_files(sample_dir), [Rule("PT Lama", "PT Baru")], apply=True)
    assert len(summary.failed_files) == 1
    assert summary.failed_files[0].path.name == "broken.xlsx"
    assert summary.total_replacements == 4


def test_run_is_idempotent(sample_dir):
    files = discover_files(sample_dir)
    rules = [Rule("PT Lama", "PT Baru")]
    first = run(files, rules, apply=True)
    second = run(files, rules, apply=True)
    assert first.total_replacements > 0
    assert second.total_replacements == 0


def test_no_backup_flag(sample_dir):
    result = process_file(
        sample_dir / "book1.xlsx", [Rule("PT Lama", "X")], apply=True, backup=False
    )
    assert result.backup is None
    assert not (sample_dir / "book1.xlsx.bak").exists()


def test_unchanged_file_is_not_rewritten(sample_dir):
    path = sample_dir / "book2.xlsx"
    before = path.stat().st_mtime_ns
    result = process_file(path, [Rule("PT Lama", "X")], apply=True)
    assert result.written is False
    assert path.stat().st_mtime_ns == before


def test_failed_save_leaves_the_original_intact(sample_dir, monkeypatch):
    import openpyxl

    path = sample_dir / "book1.xlsx"
    before = path.read_bytes()

    def boom(self, target):
        raise OSError("disk full")

    monkeypatch.setattr(openpyxl.Workbook, "save", boom)
    result = process_file(path, [Rule("PT Lama", "PT Baru")], apply=True)

    assert "disk full" in result.error
    assert path.read_bytes() == before
    assert not list(sample_dir.glob("*.emr-tmp"))
    assert not list(sample_dir.glob("*.bak"))

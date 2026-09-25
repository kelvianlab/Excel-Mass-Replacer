import openpyxl

from excel_mass_replacer.cli import main


def test_cli_dry_run_by_default(sample_dir, capsys):
    assert main(["Acme", "Globex", "-d", str(sample_dir)]) == 0
    out = capsys.readouterr().out
    assert "PREVIEW" in out and "Replacements      : 4" in out
    assert openpyxl.load_workbook(sample_dir / "book1.xlsx").active["A1"].value == (
        "Acme Holdings"
    )


def test_cli_apply(sample_dir, capsys):
    assert main(["Acme", "Globex", "-d", str(sample_dir), "--apply"]) == 0
    assert "APPLY" in capsys.readouterr().out
    assert openpyxl.load_workbook(sample_dir / "book1.xlsx").active["A1"].value == (
        "Globex Holdings"
    )


def test_cli_pairs_file(sample_dir, tmp_path, capsys):
    pairs = tmp_path / "pairs.tsv"
    pairs.write_text("# comment\nAcme\tGlobex\nlegacy\tmodern\n", encoding="utf-8")
    assert main(["-d", str(sample_dir), "--pairs-file", str(pairs), "--apply"]) == 0
    assert openpyxl.load_workbook(sample_dir / "book1.xlsx").active["A1"].value == (
        "Globex Holdings"
    )


def test_cli_missing_args(capsys):
    assert main(["-d", "."]) == 2
    assert "FIND and REPLACE" in capsys.readouterr().err


def test_cli_missing_folder(capsys):
    assert main(["a", "b", "-d", "/definitely/not/here"]) == 2
    assert "path not found" in capsys.readouterr().err


def test_cli_include_exclude(sample_dir, capsys):
    main(["Acme", "X", "-d", str(sample_dir), "--include", "book1.xlsx"])
    assert "Files scanned     : 1" in capsys.readouterr().out


def test_cli_returns_1_on_failure(sample_dir, capsys):
    (sample_dir / "broken.xlsx").write_bytes(b"not a zip")
    assert main(["Acme", "X", "-d", str(sample_dir), "--apply"]) == 1
